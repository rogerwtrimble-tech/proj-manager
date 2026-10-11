import os
import time

os.environ.setdefault("DATABASE_URL", "sqlite:///./test_proj_manager.db")
os.environ.setdefault("ADMIN_TOKEN", "test-admin")
os.environ["DATA_DIR"] = "./test_data"

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app import pipeline  # noqa: E402
from app.main import app  # noqa: E402

ADMIN = {"X-Admin-Token": "test-admin"}
DATE = "11.01.2026"


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield c


def test_pipeline_functions_demo(tmp_path=None):
    class Run:
        migration_date = DATE
        folder_url = f"Migration - {DATE}"
        log_artifact = ""
        xlsx_artifact = ""

    logs: list[str] = []
    run = Run()
    counts = pipeline.run_pipeline_steps(run, ["stage", "normalize", "precheck", "report"], logs.append)
    # Sample data deliberately contains defects:
    # duplicate address, missing price, bad price, invalid email, occupied w/o occupier
    assert counts["errors"] >= 4, logs
    assert counts["warnings"] >= 2, logs
    base = pipeline.mig_dir(DATE)
    for name in ("addresses.csv", "occupier.csv", "utilities.csv"):
        assert os.path.isfile(os.path.join(base, "combined", name)), name
    log_text = open(os.path.join(base, "reports", run.log_artifact)).read()
    assert "[ERROR]" in log_text and "[WARNING]" in log_text
    from openpyxl import load_workbook

    wb = load_workbook(os.path.join(base, "reports", run.xlsx_artifact))
    assert {"Summary", "Errors", "Warnings", "Addresses_Fixes", "Occupier_Fixes", "Utilities_Fixes"} <= set(wb.sheetnames)


def test_pipeline_api_and_script_runner(client):
    r = client.post("/api/pipeline/runs", json={"migration_date": DATE, "folder_url": ""}, headers=ADMIN)
    assert r.status_code == 201
    run_id = r.json()["id"]
    deadline = time.time() + 20
    body = {}
    while time.time() < deadline:
        body = client.get(f"/api/pipeline/runs/{run_id}").json()
        if body["status"] in ("completed", "failed"):
            break
        time.sleep(0.3)
    assert body["status"] == "completed", body
    assert body["errors"] >= 4
    r = client.get(f"/api/pipeline/runs/{run_id}/artifacts/log")
    assert r.status_code == 200 and "[ERROR]" in r.text
    r = client.get(f"/api/pipeline/runs/{run_id}/artifacts/xlsx")
    assert r.status_code == 200

    scripts = client.get("/api/scripts").json()
    assert any(s["name"] == "row_counts.py" for s in scripts)
    sid = next(s["id"] for s in scripts if s["name"] == "row_counts.py")
    r = client.post(f"/api/scripts/{sid}/run", headers=ADMIN)
    assert r.status_code == 200
    assert r.json()["last_status"] == "ok"
    assert "addresses.csv" in r.json()["last_output"]


def test_pipeline_requires_admin(client):
    r = client.post("/api/pipeline/runs", json={"migration_date": DATE})
    assert r.status_code == 401


def test_quick_run_tile_flow(client):
    r = client.post("/api/pipeline/run-quick", json={"password": "wrong", "migration_date": DATE})
    assert r.status_code == 403
    r = client.post("/api/pipeline/run-quick", json={"password": "test-admin", "migration_date": DATE})
    assert r.status_code == 202
    run_id = r.json()["run_id"]
    deadline = time.time() + 20
    body = {}
    while time.time() < deadline:
        body = client.get(f"/api/pipeline/runs/{run_id}").json()
        if body["status"] in ("completed", "failed"):
            break
        time.sleep(0.3)
    assert body["status"] == "completed", body
    status = client.get("/api/status").json()
    assert status["last_prep"] is not None
    assert status["last_prep"]["migration_date"] == DATE
    assert status["last_prep"]["errors"] >= 4
